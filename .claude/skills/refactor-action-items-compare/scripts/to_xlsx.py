# 액션 아이템 md → 엑셀(xlsx), 2026-10-06-전체.xlsx 양식.
# 아이템 하나가 한 행이 되어 담당자·상태를 적어 가며 관리할 수 있다. 시트: 액션 아이템 · 리팩토링이 아닌 발견 · 그 밖의 후보 · 점검 대상
# --compare <기준 엑셀> 을 주면 기준판(A-RMS_리팩토링_액션아이템_<날짜>.xlsx 양식)과 대조한 「<MM-DD> 비교」 시트를 더한다.
# 사용: python3 to_xlsx.py <md 파일> [<md 파일> ...] [-o 출력.xlsx] [--compare 기준.xlsx [--overrides 판정.json]]
#   - md 하나면 기본 출력은 같은 위치의 같은 이름 .xlsx
#   - md 여러 개면 한 파일로 합친다 (-o 필수). "대상" 열로 구분된다.
#   - 판정.json: {"<기준 ID>": {"verdict": "유지|변경|해소|…", "new_ids": ["<대상> A-03", …], "note": "…"}} — 원문 확인으로 정한 수동 판정
# 필요: pip install openpyxl
import sys, re, os, argparse, json, difflib
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

MODULES = ["Backend-Core", "Engine-Fire", "Middle-Proxy", "Global-Config", "Broker-Hub", "Frontend-Web", "AI"]


def clean(text):
    text = text.replace("<br>", "\n").replace("\\|", "|").replace("`", "")
    return re.sub(r"\*\*(.+?)\*\*", r"\1", text).strip()


def parse(path):
    md = open(path, encoding="utf-8").read()
    title = re.search(r"^# 리팩토링 액션 아이템 — (.+)$", md, re.M)
    meta = dict(re.findall(r"^- (점검 대상|점검일): (.+)$", md, re.M))
    doc = {"대상": title.group(1).strip() if title else os.path.basename(path),
           "점검 대상": clean(meta.get("점검 대상", "")), "점검일": meta.get("점검일", "").strip()}

    def sections(heading):
        m = re.search(r"^## " + heading + r"\n(.*?)(?=^## |\Z)", md, re.M | re.S)
        return m.group(1) if m else ""

    def blocks(body, prefix):
        out = []
        for b in re.split(r"^### ", body, flags=re.M)[1:]:
            head, _, rest = b.partition("\n")
            if not head.startswith(prefix):
                continue
            row = {"head": head.strip()}
            for label, value in re.findall(r"^\| \*\*(.+?)\*\* \| (.*) \|$", rest, re.M):
                row[label] = clean(value)
            for cap, code in re.findall(r"^\*(원문|바꾼 모습)\*.*?\n\n```[^\n]*\n(.*?)\n```", rest, re.M | re.S):
                row[cap] = code
            out.append(row)
        return out

    items = []
    for r in blocks(sections("액션 아이템"), "A-"):
        m = re.match(r"(A-\d+) \[(.+?)\] (.+?) — 우선순위 (\S+)$", r["head"])
        no, cat, name, prio = m.groups() if m else (r["head"], "", "", "")
        items.append([doc["대상"], no, cat, prio, name, r.get("위치", ""), r.get("무엇이 문제인가", ""),
                      r.get("왜 수정이 필요한가", ""), r.get("이렇게 바꾼다", ""), r.get("먼저 할 것", ""),
                      r.get("주의", ""), r.get("원문", ""), r.get("바꾼 모습", ""), "", "미착수", ""])
    findings = []
    for r in blocks(sections("리팩토링이 아닌 발견"), "F-"):
        m = re.match(r"(F-\d+) \[(.+?)\] (.+)$", r["head"])
        no, kind, name = m.groups() if m else (r["head"], "", "")
        findings.append([doc["대상"], no, kind, name, r.get("위치", ""), r.get("내용", ""), r.get("권고", ""), "", "미착수", ""])
    others = []
    for line in sections("그 밖의 후보").splitlines():
        m = re.match(r"^- (.+?) \[(.+?)\] (.+)$", line)
        if m:
            others.append([doc["대상"], clean(m.group(1)), m.group(2), clean(m.group(3))])
        elif line.startswith("- "):   # 폴더 순위 요약처럼 [카테고리] 형식이 아닌 줄도 빠뜨리지 않는다
            others.append([doc["대상"], "", "", clean(line[2:])])
    targets = [[doc["대상"], doc["점검 대상"], doc["점검일"], len(items), len(findings)]]
    return items, findings, others, targets


STATUS = '"미착수,진행 중,완료,보류,하지 않음"'
SHEETS = [
    ("액션 아이템", ["대상", "#", "카테고리", "우선순위", "제목", "위치", "무엇이 문제인가", "왜 수정이 필요한가",
                  "이렇게 바꾼다", "먼저 할 것", "주의", "원문", "바꾼 모습", "담당자", "상태", "비고"],
     [18, 7, 16, 9, 34, 38, 50, 44, 50, 38, 34, 50, 50, 10, 10, 20], "O"),
    ("리팩토링이 아닌 발견", ["대상", "#", "구분", "제목", "위치", "내용", "권고", "담당자", "상태", "비고"],
     [18, 7, 12, 36, 38, 60, 44, 10, 10, 20], "I"),
    ("그 밖의 후보", ["대상", "위치", "카테고리", "내용"], [18, 40, 14, 70], None),
    ("점검 대상", ["대상", "점검 대상", "점검일", "액션 아이템", "리팩토링이 아닌 발견"], [24, 70, 12, 12, 18], None),
]
PRIO_FILL = {"높음": "F8D7DA", "중간": "FFF3CD", "낮음": "E2EFDA"}
VERDICT_FILL = {"유지": "E2EFDA", "변경": "FFF3CD", "신규": "F8D7DA", "해소": "D9EAD3", "확인 필요": "FFFF99"}
THIN = Side(style="thin", color="C9D3D8")


def style_sheet(ws, header, widths, nrows):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0E4F63")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in ws.iter_rows(min_row=2, max_row=nrows + 1, max_col=len(header)):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
            if header[cell.column - 1] in ("원문", "바꾼 모습"):
                cell.font = Font(name="Menlo", size=9)
            if header[cell.column - 1] == "우선순위" and cell.value in PRIO_FILL:
                cell.fill = PatternFill("solid", fgColor=PRIO_FILL[cell.value])
            if header[cell.column - 1] == "판정" and cell.value in VERDICT_FILL:
                cell.fill = PatternFill("solid", fgColor=VERDICT_FILL[cell.value])


def write(rows_by_sheet, out, compare_rows=None, compare_name=None):
    wb = Workbook(); wb.remove(wb.active)
    for (name, header, widths, status_col), rows in zip(SHEETS, rows_by_sheet):
        ws = wb.create_sheet(name)
        ws.append(header)
        for row in rows:
            ws.append(row)
        style_sheet(ws, header, widths, len(rows))
        ws.freeze_panes = "C2"
        if rows:
            ws.auto_filter.ref = ws.dimensions
        if status_col and rows:
            dv = DataValidation(type="list", formula1=STATUS, allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"{status_col}2:{status_col}{len(rows) + 1}")
    if compare_rows is not None:
        header = ["모듈", "구분", "기준 ID", "기준 우선순위", "기준 제목", "기준 위치", "새 대상·#", "새 제목", "판정", "비고"]
        ws = wb.create_sheet(compare_name)
        ws.append(header)
        for row in compare_rows:
            ws.append(row)
        style_sheet(ws, header, [14, 8, 10, 10, 44, 46, 26, 44, 11, 60], len(compare_rows))
        ws.freeze_panes = "C2"
        if compare_rows:
            ws.auto_filter.ref = f"A1:J{len(compare_rows) + 1}"
        # 판정 집계 (오른쪽, 수식)
        verdicts = ["유지", "변경", "해소", "확인 필요", "신규"]
        for j, v in enumerate(["판정", "아이템", "발견"], 12):
            c = ws.cell(1, j, v)
            c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="0E4F63")
        ws.column_dimensions["L"].width = 12
        for i, v in enumerate(verdicts, 2):
            ws.cell(i, 12, v)
            ws.cell(i, 13, f'=COUNTIFS($I:$I,$L{i},$B:$B,"아이템")')
            ws.cell(i, 14, f'=COUNTIFS($I:$I,$L{i},$B:$B,"발견")')
        ws.cell(2 + len(verdicts), 12, "합계")
        ws.cell(2 + len(verdicts), 13, f"=SUM(M2:M{1 + len(verdicts)})")
        ws.cell(2 + len(verdicts), 14, f"=SUM(N2:N{1 + len(verdicts)})")
    wb.save(out)


# ──────────────────────────── 기준판 비교
def module_of(target):
    head = target.split()[0] if target else ""
    return head if head in MODULES else next((m for m in MODULES if m in target), head)


def loc_keys(text):
    """위치 칸에서 (모듈 기준 상대 경로, 줄) 목록. 경로 앞의 'Java-Service-Tree-Framework-<모듈>/' 는 떼어 낸다."""
    keys = []
    for p, l in re.findall(r"((?:Java-Service-Tree-Framework-[\w-]+/)?src/[^\s:`]+?):(\d+)", text or ""):
        keys.append((re.sub(r"^Java-Service-Tree-Framework-[\w-]+/", "", p), int(l)))
    return keys


def nums(v):
    return [int(x) for x in re.findall(r"\d+", str(v or ""))]


def sim(a, b):
    return difflib.SequenceMatcher(None, a or "", b or "").ratio()


def compare(base_path, items, findings, overrides, tol=5):
    base = load_workbook(base_path, read_only=True)
    mods = {module_of(r[0]) for r in items + findings}
    new_i = {f"{r[0]} {r[1]}": (module_of(r[0]), r[4], loc_keys(r[5])) for r in items}
    new_f = {f"{r[0]} {r[1]}": (module_of(r[0]), r[3], loc_keys(r[4])) for r in findings}
    blocs = {}
    for r in base["위치상세"].iter_rows(min_row=2, values_only=True):
        if r[0] and r[1] in mods:
            blocs.setdefault(r[0], []).append((r[4], (nums(r[5]) or [None])[0]))

    def hits(mod, locs, pool):
        out = set()
        for f, l in locs:
            for k, (m, _, ks) in pool.items():
                if m == mod and any(kf == f and l is not None and abs(kl - l) <= tol for kf, kl in ks):
                    out.add(k)
        return out

    rows, used = [], set()
    targets = []
    for r in base["액션아이템"].iter_rows(min_row=2, values_only=True):
        if r[1] and r[2] in mods:
            targets.append(("아이템", r[2], r[1], r[3], r[5], r[6], r[7], blocs.get(r[1]) or [(r[6], (nums(r[7]) or [None])[0])], new_i))
    for r in base["리팩토링외발견"].iter_rows(min_row=2, values_only=True):
        if r[0] and r[2] in mods:
            targets.append(("발견", r[2], r[0], r[1], r[3], r[4], r[5], [(r[4], (nums(r[5]) or [None])[0])], new_f))
    for kind, mod, oid, prio, title, f, l, locs, pool in targets:
        cand = [k for k, (m, _, _) in pool.items() if m == mod]
        best = max(cand, key=lambda k: sim(title, pool[k][1]), default=None)
        near = hits(mod, locs, pool)
        notes = []
        if best and sim(title, pool[best][1]) >= 0.6:
            ids = [best]
            notes.append("제목·내용 일치")
            if near - {best}:
                notes.append("같은 위치의 다른 아이템(원인 다름): " + ", ".join(sorted(near - {best})))
        else:
            ids = sorted(k for k in near if sim(title, pool[k][1]) >= 0.3)
            if ids:
                notes.append("위치 일치(같은 파일·줄 ±%d), 제목은 다시 씀" % tol)
            if near - set(ids):
                notes.append("위치만 인접·원인 다름: " + ", ".join(sorted(near - set(ids))))
        o = overrides.get(oid, {})
        ids = o.get("new_ids", ids)
        verdict = o.get("verdict", "유지" if ids else "확인 필요")
        note = o.get("note", " · ".join(notes) or "같은 위치·제목의 새 항목 없음")
        used |= set(ids)
        rows.append([mod, kind, oid, prio, title, f"{f}:{l}", ", ".join(ids) or None,
                     " / ".join(pool[i][1] for i in ids if i in pool) or None, verdict, note])
    for k, (m, t, _) in list(new_i.items()):
        if k not in used:
            rows.append([m, "아이템", None, None, None, None, k, t, "신규", None])
    for k, (m, t, _) in list(new_f.items()):
        if k not in used:
            rows.append([m, "발견", None, None, None, None, k, t, "신규", None])
    order = {m: i for i, m in enumerate(MODULES)}
    rows.sort(key=lambda x: (order.get(x[0], 99), x[1] != "아이템", x[2] is None, x[2] or "", x[6] or ""))
    return rows


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("md", nargs="+")
    ap.add_argument("-o", "--out")
    ap.add_argument("--compare", help="대조할 기준 엑셀(A-RMS_리팩토링_액션아이템_<날짜>.xlsx 양식)")
    ap.add_argument("--overrides", help="수동 판정 JSON")
    a = ap.parse_args()
    if len(a.md) > 1 and not a.out:
        sys.exit("md 파일이 여러 개면 -o 로 출력 파일을 지정하세요.")
    out = a.out or a.md[0][:-3] + ".xlsx"
    merged = [[], [], [], []]
    for path in a.md:
        for acc, part in zip(merged, parse(path)):
            acc.extend(part)
    crows = cname = None
    if a.compare:
        ov = json.load(open(a.overrides, encoding="utf-8")) if a.overrides else {}
        crows = compare(a.compare, merged[0], merged[1], ov)
        d = re.search(r"(\d{4})-(\d{2})-(\d{2})", os.path.basename(a.compare))
        cname = f"{d.group(2)}-{d.group(3)} 비교" if d else "기준판 비교"
    write(merged, out, crows, cname)
    print(f"{out}  (액션 아이템 {len(merged[0])}건, 리팩토링이 아닌 발견 {len(merged[1])}건, 그 밖의 후보 {len(merged[2])}건)")
    if crows is not None:
        cnt = {}
        for r in crows:
            cnt[(r[1], r[8])] = cnt.get((r[1], r[8]), 0) + 1
        print(f"  {cname}: {len(crows)}행 {cnt}")
        todo = [r[2] for r in crows if r[8] == "확인 필요"]
        if todo:
            print("  확인 필요(원문 확인 후 --overrides 로 판정):", ", ".join(todo))


if __name__ == "__main__":
    main()
