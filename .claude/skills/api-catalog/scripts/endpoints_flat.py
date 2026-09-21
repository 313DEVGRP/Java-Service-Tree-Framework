# -*- coding: utf-8 -*-
"""엔드포인트 플랫 CSV + 컨트롤러별 시트 xlsx.

expand_inherit.py 산출물을 사람이 훑는 형태로 편다.
  입력 : $CATALOG_TASK/sources/endpoints_expanded.csv · inventory.json
  출력 : $CATALOG_TASK/artifacts/<repo>_endpoints_flat.csv
         $CATALOG_TASK/artifacts/<repo>_endpoints_by-controller.xlsx

자체 선언 = 1건 1행. 상속분 = 부모 매핑 1행 + 상속 컨트롤러 열거(컨트롤러별 반복 없음).
provider 컬럼: TreeAbstractController·TreeMapAbstractController 상속분은 TreeFramework 가
공통 제공하는 트리 CRUD 라 프론트가 호출하지 않는 것이 정상일 수 있다 — 그 구분을 값으로 남긴다.
"""
import csv, io, json, os, re, collections
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

_REPO  = os.environ.get('CATALOG_REPO', 'Java-Service-Tree-Framework-Backend-Core')
_SHORT = _REPO.replace('Java-Service-Tree-Framework-', '').lower()
TASK = os.environ.get('CATALOG_TASK', 'tasks/api-catalog')
SRC  = os.path.join(TASK, 'sources')
OUT_CSV  = os.path.join(TASK, 'artifacts', '%s_endpoints_flat.csv' % _SHORT)
OUT_XLSX = os.path.join(TASK, 'artifacts', '%s_endpoints_by-controller.xlsx' % _SHORT)

HEADER = ['no', 'kind', 'provider', 'controllerClass', 'package', 'endpoint', 'httpMethod',
          'path', 'basePath', 'javaMethod', 'sourceFile', 'line',
          'inheritedByCount', 'inheritedBy', 'note']

pkg_of = lambda f: f.rsplit('/', 1)[0].replace('/', '.') if '/' in f else ''

rows = list(csv.DictReader(io.open(os.path.join(SRC, 'endpoints_expanded.csv'), encoding='utf-8-sig')))
inv  = json.load(io.open(os.path.join(SRC, 'inventory.json'), encoding='utf-8'))['perController']
base_of = lambda c: inv.get(c, {}).get('base', '')

own_rows  = [r for r in rows if r['출처'] == '자체']
# 재정의 판정은 메서드명이 아니라 (컨트롤러, HTTP, 전체경로) 로 한다.
own_paths = set((r['controller'], r['HTTP'], r['path']) for r in own_rows)

out = []

# ── 1) 자체 선언: 1건 1행 ──────────────────────────────────────────────
for r in sorted(own_rows, key=lambda x: (x['controller'], x['path'], x['HTTP'])):
    out.append(['', 'own', '자체 선언', r['controller'], pkg_of(r['선언파일']),
                '%s %s' % (r['HTTP'], r['path']), r['HTTP'], r['path'],
                base_of(r['controller']), r['javaMethod'],
                r['선언파일'], r['line'], '', '', ''])

# ── 2) 상속분: 부모 매핑 1행 + 상속 컨트롤러 열거 ──────────────────────
inh = {}
for r in rows:
    if r['출처'] == '자체':
        continue
    parent = r['출처'][r['출처'].index('(') + 1:-1]
    b   = base_of(r['controller'])
    rel = r['path'][len(b):] if b and r['path'].startswith(b) else r['path']
    e = inh.setdefault((parent, r['HTTP'], rel, r['javaMethod']),
                       {'ctls': set(), 'file': r['선언파일'], 'line': r['line']})
    e['ctls'].add(r['controller'])

for (parent, http, rel, jm), e in sorted(inh.items(), key=lambda kv: (kv[0][0], kv[0][2], kv[0][1])):
    ctls = sorted(e['ctls'])
    ov   = [c for c in ctls if (c, http, base_of(c) + rel) in own_paths]
    out.append(['', 'inherited', '상속(%s)' % parent, parent, pkg_of(e['file']),
                '%s {basePath}%s' % (http, rel), http, '{basePath}%s' % rel, '',
                jm, e['file'], e['line'], len(ctls),
                ';'.join('%s@%s' % (c, base_of(c)) for c in ctls),
                ('자체 선언으로 재정의: ' + ';'.join(ov)) if ov else ''])

for i, r in enumerate(out, 1):
    r[0] = i

os.makedirs(os.path.dirname(OUT_CSV), exist_ok=True)
with io.open(OUT_CSV, 'w', encoding='utf-8-sig', newline='') as f:
    w = csv.writer(f)
    w.writerow(HEADER)
    w.writerows(out)

# ── 3) 컨트롤러별 시트 xlsx ────────────────────────────────────────────
BAD  = re.compile(r'[\[\]:*?/\\]')
GREY = PatternFill('solid', fgColor='EFEFEF')
dicts = [dict(zip(HEADER, r)) for r in out]

def sheet_name(v, used):
    n = BAD.sub('-', str(v)).strip()[:31] or 'EMPTY'
    if n in used:
        for i in range(2, 100):
            c = n[:31 - len('~%d' % i)] + '~%d' % i
            if c not in used:
                n = c; break
    used.add(n)
    return n

wb = Workbook(); idx = wb.active; idx.title = '_목차'
smy = wb.create_sheet('_요약')
provs = sorted({d['provider'] for d in dicts})
smy.append(['kind'] + provs + ['합계'])
for kv in sorted({d['kind'] for d in dicts}):
    sub = [d for d in dicts if d['kind'] == kv]
    smy.append([kv] + [sum(1 for d in sub if d['provider'] == p) for p in provs] + [len(sub)])
smy.append(['합계'] + [sum(1 for d in dicts if d['provider'] == p) for p in provs] + [len(dicts)])
smy['A%d' % (smy.max_row + 2)] = ('상속(...) = TreeAbstractController·TreeMapAbstractController 에서 물려받아 생긴 URL. '
                                  'TreeFramework 가 공통 제공하는 트리 CRUD 라 프론트에서 호출하지 않는 것이 정상일 수 있다.')
for c in range(1, len(provs) + 3):
    smy.cell(1, c).font = Font(bold=True)
    smy.column_dimensions[get_column_letter(c)].width = 42 if c > 1 else 14

groups = collections.OrderedDict()
for d in dicts:
    groups.setdefault(d['controllerClass'], []).append(d)
idx.append(['#', 'controllerClass', '시트명', '행수', '상속', '자체 선언'])
used = {'_목차', '_요약'}
for i, (v, rs) in enumerate(groups.items(), 1):
    sn = sheet_name(v, used)
    fw = sum(1 for d in rs if d['provider'].startswith('상속'))
    idx.append([i, v, sn, len(rs), fw, len(rs) - fw])
    ws = wb.create_sheet(sn)
    ws.append(HEADER)
    for d in rs:
        ws.append([d[c] for c in HEADER])
    for c in range(1, len(HEADER) + 1):
        ws.cell(1, c).font = Font(bold=True)
    for n, d in enumerate(rs, 2):
        if d['provider'].startswith('상속'):
            for c in range(1, len(HEADER) + 1):
                ws.cell(n, c).fill = GREY
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = 'A1:%s%d' % (get_column_letter(len(HEADER)), len(rs) + 1)
    for c, cn in enumerate(HEADER, 1):
        w_ = max([len(cn)] + [len(str(d[cn])) for d in rs])
        ws.column_dimensions[get_column_letter(c)].width = min(max(w_ + 2, 8), 60)
for c, w_ in ((1, 5), (2, 40), (3, 34), (4, 8), (5, 16), (6, 12)):
    idx.column_dimensions[get_column_letter(c)].width = w_
for c in range(1, 7):
    idx.cell(1, c).font = Font(bold=True)
idx.freeze_panes = 'A2'; idx.auto_filter.ref = 'A1:F%d' % (len(groups) + 1)
wb.save(OUT_XLSX)

n_own = sum(1 for d in dicts if d['kind'] == 'own')
print('%s\n%s' % (OUT_CSV, OUT_XLSX))
print('총 %d행 = 자체 %d + 상속 부모매핑 %d | 상속 표기 %d | 상속 전개 실인스턴스 %d | 도달 가능 합계 %d | 시트 %d'
      % (len(dicts), n_own, len(dicts) - n_own,
         sum(1 for d in dicts if d['provider'].startswith('상속')),
         sum(int(d['inheritedByCount']) for d in dicts if d['kind'] == 'inherited'),
         n_own + sum(int(d['inheritedByCount']) for d in dicts if d['kind'] == 'inherited'),
         len(groups)))
