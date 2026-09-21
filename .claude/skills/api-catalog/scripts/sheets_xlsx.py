# -*- coding: utf-8 -*-
"""CSV → 키 컬럼별 시트로 나눈 xlsx. 원본 CSV 는 읽기만 한다.

사용: python sheets_xlsx.py <csv> <키컬럼|-> <out.xlsx> [요약컬럼]
  키컬럼에 `-` 를 주면 시트를 나누지 않고 한 장(`data`)에 그 CSV 순서대로 담는다.
  요약컬럼을 주면 `_요약` 시트에 (요약컬럼 × provider) 교차 집계를 만든다.
  provider 컬럼이 있으면 `상속(...)` 행을 회색으로 깐다 — TreeFramework 상속분 구분용.
엑셀 시트명 제약(31자·`[]:*?/\\` 금지)은 자동 치환하고 원래 값은 `_목차` 에 남긴다.
"""
import csv, io, os, re, sys, collections
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

BAD  = re.compile(r'[\[\]:*?/\\]')
GREY = PatternFill('solid', fgColor='EFEFEF')
NOTE = ('상속(...) = TreeAbstractController·TreeMapAbstractController 에서 물려받아 생긴 URL. '
        'TreeFramework 가 공통 제공하는 트리 CRUD 라 프론트에서 호출하지 않는 것이 정상일 수 있다.')

def sheet_name(v, used):
    n = BAD.sub('-', str(v)).strip()[:31] or 'EMPTY'
    if n in used:
        for i in range(2, 100):
            c = n[:31 - len('~%d' % i)] + '~%d' % i
            if c not in used:
                n = c; break
    used.add(n)
    return n

def write_sheet(wb, name, cols, rows, inh):
    ws = wb.create_sheet(name)
    ws.append(cols)
    for r in rows:
        ws.append([r[c] for c in cols])
    for c in range(1, len(cols) + 1):
        ws.cell(1, c).font = Font(bold=True)
    for n, r in enumerate(rows, 2):
        if inh(r):
            for c in range(1, len(cols) + 1):
                ws.cell(n, c).fill = GREY
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = 'A1:%s%d' % (get_column_letter(len(cols)), len(rows) + 1)
    for c, cn in enumerate(cols, 1):
        w = max([len(cn)] + [len(str(r[cn])) for r in rows])
        ws.column_dimensions[get_column_letter(c)].width = min(max(w + 2, 8), 60)
    return ws

def main():
    if len(sys.argv) < 4:
        print(__doc__); return 64
    src, key, out = sys.argv[1], sys.argv[2], sys.argv[3]
    sumcol = sys.argv[4] if len(sys.argv) > 4 else None

    rows = list(csv.DictReader(io.open(src, encoding='utf-8-sig')))
    if not rows:
        print('빈 CSV: %s' % src); return 1
    cols = list(rows[0].keys())
    has_prov = 'provider' in cols
    inh = lambda r: has_prov and r['provider'].startswith('상속')

    wb = Workbook(); idx = wb.active; idx.title = '_목차'
    used = {'_목차'}

    if sumcol and has_prov:
        used.add('_요약')
        smy = wb.create_sheet('_요약')
        provs = sorted({r['provider'] for r in rows})
        smy.append([sumcol] + provs + ['합계'])
        for kv in sorted({r[sumcol] for r in rows}):
            sub = [r for r in rows if r[sumcol] == kv]
            smy.append([kv] + [sum(1 for r in sub if r['provider'] == p) for p in provs] + [len(sub)])
        smy.append(['합계'] + [sum(1 for r in rows if r['provider'] == p) for p in provs] + [len(rows)])
        smy['A%d' % (smy.max_row + 2)] = NOTE
        for c in range(1, len(provs) + 3):
            smy.cell(1, c).font = Font(bold=True)
            smy.column_dimensions[get_column_letter(c)].width = 40 if c > 1 else 16

    if key == '-':                      # 시트 분할 없음 — 한 장에 전부
        wb.remove(idx)
        write_sheet(wb, 'data', cols, rows, inh)
        os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)
        wb.save(out)
        print('%s\n단일 시트 | %d행%s'
              % (out, len(rows),
                 (' | 상속 %d' % sum(1 for r in rows if inh(r))) if has_prov else ''))
        return 0

    groups = collections.OrderedDict()
    for r in rows:
        groups.setdefault(r[key], []).append(r)

    head = ['#', key, '시트명', '행수'] + (['상속', '그 외'] if has_prov else [])
    idx.append(head)
    for i, (v, rs) in enumerate(groups.items(), 1):
        sn = sheet_name(v, used)
        n_inh = sum(1 for r in rs if inh(r))
        idx.append([i, v, sn, len(rs)] + ([n_inh, len(rs) - n_inh] if has_prov else []))
        write_sheet(wb, sn, cols, rs, inh)
    for c, w in enumerate((5, 46, 34, 8, 10, 10)[:len(head)], 1):
        idx.column_dimensions[get_column_letter(c)].width = w
    for c in range(1, len(head) + 1):
        idx.cell(1, c).font = Font(bold=True)
    idx.freeze_panes = 'A2'
    idx.auto_filter.ref = 'A1:%s%d' % (get_column_letter(len(head)), len(groups) + 1)

    os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)
    wb.save(out)
    print('%s\n시트 %d (+목차%s) | %d행%s'
          % (out, len(groups), '·요약' if (sumcol and has_prov) else '', len(rows),
             (' | 상속 %d' % sum(1 for r in rows if inh(r))) if has_prov else ''))
    return 0

if __name__ == '__main__':
    sys.exit(main())
