# 액션 아이템 md → 엑셀(xlsx). 팀 템플릿(assets/action-items-template.xlsx, A-RMS_리팩토링_액션아이템 양식)의
# 시트·열·서식을 그대로 쓰고 데이터만 채운다.
#   시트: 요약 · 액션아이템 · 위치상세 · 리팩토링외발견 · 그밖의후보 · 리뷰반영 · 카테고리기준
# 사용: python3 to_xlsx.py <md 파일> [<md 파일> ...] [-o 출력.xlsx]
#   - md 하나면 기본 출력은 같은 위치의 같은 이름 .xlsx
#   - md 여러 개면 한 파일로 합친다 (-o 필수). "모듈" 열로 구분된다.
# 필요: pip install openpyxl
import argparse, copy, os, re, sys
import openpyxl
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter

TEMPLATE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "assets", "action-items-template.xlsx")
MODULES = {  # 모듈 이름 → (ID 접두어, 스택)
    "Backend-Core": ("BC", "Java 11 · Spring Boot 2.6 · Hibernate/JPA · MySQL"),
    "Engine-Fire": ("EF", "Java 21 · Spring Boot 3.5 · OpenSearch"),
    "Middle-Proxy": ("MP", "Java 11 · Spring Boot 2.6 · Spring Cloud Gateway(WebFlux)"),
    "AI": ("AI", "Java 21 · Spring Boot 3.5 · Spring AI · WebFlux"),
    "Broker-Hub": ("BH", "Java 21 · Spring Boot 3.5 · STOMP · Redis"),
    "Global-Config": ("GC", "Java 11 · Spring Boot 2.6 · Config Server(MVC)"),
    "Frontend-Web": ("FE", "jQuery · Bootstrap 3 (arms/ · backoffice/)"),
}
PRIO = {"높음": 0, "중간": 1, "낮음": 2}
PRIO_FILL = {"높음": "F8D7D3", "중간": "FCEFD2", "낮음": "E3EEF1"}


def clean(text):
    text = text.replace("<br>", "\n").replace("\\|", "|").replace("`", "")
    return re.sub(r"\*\*(.+?)\*\*", r"\1", text).strip()


def module_of(title, target):
    for name in sorted(MODULES, key=len, reverse=True):
        if f"Java-Service-Tree-Framework-{name}" in target or title.split()[0] == name:
            return name
    return title.split()[0] if title else "?"


def parse(path):
    md = open(path, encoding="utf-8").read()
    title = (re.search(r"^# 리팩토링 액션 아이템 — (.+)$", md, re.M) or [None, os.path.basename(path)])[1].strip()
    meta = dict(re.findall(r"^- (점검 대상|점검일|기준 브랜치): (.+)$", md, re.M))
    head = re.search(r"\(([0-9a-f]{6,})\)", meta.get("기준 브랜치", ""))
    doc = {"제목": title, "모듈": module_of(title, meta.get("점검 대상", "")), "HEAD": head.group(1) if head else "",
           "점검일": meta.get("점검일", "").strip()}

    def section(heading):
        m = re.search(r"^## " + heading + r"\n(.*?)(?=^## |\Z)", md, re.M | re.S)
        return m.group(1) if m else ""

    def blocks(body, prefix):
        out = []
        for b in re.split(r"^### ", body, flags=re.M)[1:]:
            head_line, _, rest = b.partition("\n")
            if head_line.startswith(prefix):
                row = {"head": head_line.strip()}
                row.update({k: v for k, v in re.findall(r"^\| \*\*(.+?)\*\* \| (.*) \|$", rest, re.M)})
                out.append(row)
        return out

    def locations(cell):
        locs = []
        for part in cell.split("<br>"):
            m = re.match(r"\s*•?\s*`([^`]+?):(\d+)`\s*(.*)$", part)
            if m:
                locs.append((m.group(1), int(m.group(2)), clean(m.group(3))))
            elif clean(part.lstrip("• ")):
                locs.append((clean(part.lstrip("• ")), None, ""))
        return locs or [("", None, "")]

    items = []
    for r in blocks(section("액션 아이템"), "A-"):
        m = re.match(r"A-(\d+) \[(.+?)\] (.+?) — 우선순위 (\S+)$", r["head"])
        num, cat, name, prio = m.groups() if m else ("0", "", r["head"], "")
        items.append(dict(num=int(num), cat=cat, title=clean(name), prio=prio, locs=locations(r.get("위치", "")),
                          what=clean(r.get("무엇이 문제인가", "")), why=clean(r.get("왜 수정이 필요한가", "")),
                          how=clean(r.get("이렇게 바꾼다", "")), first=clean(r.get("먼저 할 것", "")),
                          caution=clean(r.get("주의", "")), size=clean(r.get("난이도", "")),
                          stage=clean(r.get("권장 단계", "")), pre=clean(r.get("선행 아이템", "")),
                          split=clean(r.get("동작 변경 분리", ""))))
    findings = []
    for r in blocks(section("리팩토링이 아닌 발견"), "F-"):
        m = re.match(r"F-(\d+) \[(.+?)\] (.+)$", r["head"])
        num, kind, name = m.groups() if m else ("0", "", r["head"])
        f, l, _ = locations(r.get("위치", ""))[0]
        findings.append(dict(num=int(num), kind=kind, title=clean(name), file=f, line=l,
                             body=clean(r.get("내용", "")), rec=clean(r.get("권고", ""))))
    others = []
    for line in section("그 밖의 후보").splitlines():
        m = re.match(r"^- `([^`]+?):(\d+)` \[(.+?)\] (.+)$", line)
        if m:
            others.append([m.group(3), m.group(1), int(m.group(2)), clean(m.group(4))])
        elif line.startswith("- "):  # 위치·카테고리 형식이 아닌 줄도 빠뜨리지 않는다
            others.append(["", "", None, clean(line[2:])])
    reviews = []
    for line in section("리뷰 반영").splitlines():
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        if line.startswith("|") and len(cells) == 5 and cells[0] not in ("리뷰어", "---") and not set(cells[0]) <= set("-: "):
            reviews.append([clean(c) for c in cells])
    return doc, items, findings, others, reviews


def styles(ws, row, ncol):
    return [copy.copy(ws.cell(row, c)._style) for c in range(1, ncol + 1)]


def fill_sheet(ws, rows, prio_col=None):
    st = styles(ws, 2, ws.max_column)
    ws.delete_rows(2, ws.max_row)
    for i, row in enumerate(rows, 2):
        for j, v in enumerate(row, 1):
            c = ws.cell(i, j, v)
            c._style = copy.copy(st[j - 1])
            if j == prio_col and v in PRIO_FILL:
                c.fill = PatternFill("solid", fgColor=PRIO_FILL[v])
    ws.freeze_panes = "A2"
    # 템플릿의 데이터 시트는 엑셀 표(Table)라 필터가 표에 들어 있다. 표 범위를 데이터 크기로 맞추고
    # 시트 단위 자동 필터는 걸지 않는다(표와 겹치면 엑셀이 '복구' 를 묻는다). 빈 표도 머리글+1행은 있어야 한다.
    ref = f"A1:{get_column_letter(ws.max_column)}{max(len(rows), 1) + 1}"
    for table in ws.tables.values():
        table.ref = ref
        if table.autoFilter is not None:
            table.autoFilter.ref = ref


def build(docs, out):
    wb = openpyxl.load_workbook(TEMPLATE)
    # 같은 모듈 md 가 여럿이면(패키지 단위로 나눈 경우) ID 접두어에 순번을 붙인다: BC1-01, BC2-01 …
    per_mod = {}
    for d in docs:
        per_mod.setdefault(d[0]["모듈"], []).append(d)
    item_rows, loc_rows, find_rows, other_rows, review_rows, mods = [], [], [], [], [], []
    for mod, ds in per_mod.items():
        mods.append((mod, ds[0][0]["HEAD"]))
        base = MODULES.get(mod, (mod[:2].upper(), ""))[0]
        for k, (doc, items, findings, others, reviews) in enumerate(ds, 1):
            pre = base if len(ds) == 1 else f"{base}{k}"
            for it in sorted(items, key=lambda x: (PRIO.get(x["prio"], 9), x["num"])):
                iid = f"{pre}-{it['num']:02d}"
                f, l, meth = it["locs"][0]
                extra = "\n".join(f"• {a}:{b} {c}".rstrip() if b else f"• {a}" for a, b, c in it["locs"][1:]) or None
                item_rows.append([None, iid, mod, it["prio"], it["cat"], it["title"], f, l, meth or None, extra,
                                  it["what"], it["why"], it["how"], it["first"] or None, it["caution"] or None,
                                  it["size"] or None, doc["HEAD"], it["cat"].split(" ")[0], it["stage"] or None,
                                  it["pre"] or None, it["split"] or None])
                loc_rows += [[iid, mod, it["prio"], it["cat"], a, b, c or None, it["title"]] for a, b, c in it["locs"]]
            find_rows += [[f"{pre}-F{x['num']:02d}", x["kind"], mod, x["title"], x["file"], x["line"], x["body"], x["rec"]]
                          for x in sorted(findings, key=lambda x: x["num"])]
            other_rows += [[mod] + o for o in others]
            review_rows += reviews
    item_rows.sort(key=lambda r: ([m for m, _ in mods].index(r[2]), PRIO.get(r[3], 9)))
    for n, r in enumerate(item_rows, 1):
        r[0] = n

    fill_sheet(wb["액션아이템"], item_rows, prio_col=4)
    fill_sheet(wb["위치상세"], loc_rows, prio_col=3)
    fill_sheet(wb["리팩토링외발견"], find_rows)
    fill_sheet(wb["그밖의후보"], other_rows)
    fill_sheet(wb["리뷰반영"], review_rows)
    write_summary(wb["요약"], mods, docs, item_rows)
    wb.save(out)
    return len(item_rows), len(loc_rows), len(find_rows), len(other_rows), len(review_rows)


def write_summary(ws, mods, docs, item_rows):
    s_title, s_sub, s_note = copy.copy(ws["A1"]._style), copy.copy(ws["A2"]._style), copy.copy(ws["A41"]._style)
    s_head, s_body, s_bold = styles(ws, 4, 9), styles(ws, 5, 9), styles(ws, 12, 9)
    ws.delete_rows(1, ws.max_row)
    dates = sorted({d[0]["점검일"] for d in docs if d[0]["점검일"]})
    ws["A1"], ws["A1"]._style = "A-RMS 리팩토링 액션 아이템 — 모듈 요약", s_title
    ws["A2"] = (f"점검일 {', '.join(dates)} · 기준: 리팩토링 가이드북(자바 RF-00~12·VO / 자바스크립트 JS-00~13) · "
                "숫자는 '액션아이템'·'리팩토링외발견'·'그밖의후보' 시트를 세는 수식")
    ws["A2"]._style = s_sub
    for j, v in enumerate(["모듈", "스택", "HEAD(dev)", "액션 아이템", "높음", "중간", "낮음", "리팩토링 외 발견", "그 밖의 후보"], 1):
        ws.cell(4, j, v)._style = copy.copy(s_head[j - 1])
    for i, (mod, head) in enumerate(mods):
        r = 5 + i
        vals = [mod, MODULES.get(mod, ("", ""))[1], head, f"=COUNTIFS(액션아이템!$C:$C,$A{r})"]
        vals += [f"=COUNTIFS(액션아이템!$C:$C,$A{r},액션아이템!$D:$D,{c}$4)" for c in "EFG"]
        vals += [f"=COUNTIFS(리팩토링외발견!$C:$C,$A{r})", f"=COUNTIFS(그밖의후보!$A:$A,$A{r})"]
        for j, v in enumerate(vals, 1):
            ws.cell(r, j, v)._style = copy.copy(s_body[j - 1])
    tr = 5 + len(mods)
    ws.cell(tr, 1, "합계")._style = copy.copy(s_bold[0])
    for j in range(4, 10):
        col = get_column_letter(j)
        ws.cell(tr, j, f"=SUM({col}5:{col}{tr - 1})")._style = copy.copy(s_bold[j - 1])
    ws.cell(tr + 2, 1, "카테고리별 분포 (액션 아이템 수)")._style = s_title
    hr = tr + 3
    names = [m for m, _ in mods]
    for j, v in enumerate(["카테고리"] + names + ["합계"], 1):
        ws.cell(hr, j, v)._style = copy.copy(s_head[0])
    cats = sorted({r[17] for r in item_rows}, key=lambda c: (c.startswith("JS"), c == "VO", c))
    last = get_column_letter(1 + len(names))
    for i, cat in enumerate(cats, 1):
        r = hr + i
        ws.cell(r, 1, cat)._style = copy.copy(s_body[0])
        for j in range(2, 2 + len(names)):
            col = get_column_letter(j)
            ws.cell(r, j, f"=COUNTIFS(액션아이템!$C:$C,{col}${hr},액션아이템!$R:$R,$A{r})")._style = copy.copy(s_body[3])
        ws.cell(r, 2 + len(names), f"=SUM(B{r}:{last}{r})")._style = copy.copy(s_body[3])
    ws.cell(hr + len(cats) + 2, 1, "※ 수식은 Excel 에서 열 때 계산된다(파일 생성 환경에 계산 엔진 없음).")._style = s_note


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("md", nargs="+")
    ap.add_argument("-o", "--out")
    a = ap.parse_args()
    if len(a.md) > 1 and not a.out:
        sys.exit("md 파일이 여러 개면 -o 로 출력 파일을 지정하세요.")
    out = a.out or a.md[0][:-3] + ".xlsx"
    n = build([parse(p) for p in a.md], out)
    print(f"{out}  (액션 아이템 {n[0]}건, 위치 {n[1]}곳, 리팩토링이 아닌 발견 {n[2]}건, 그 밖의 후보 {n[3]}건, 리뷰 반영 {n[4]}건)")


if __name__ == "__main__":
    main()
